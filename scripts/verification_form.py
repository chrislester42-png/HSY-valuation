#!/usr/bin/env python3
"""
verification_form.py

Make the Word form a person uses to check the workbook's actual-year figures against
the filing. Every row is already filled in from the workbook: row name, cell, year,
and the value. The person types only the page where they found it and Y or N.

    python3 scripts/verification_form.py
    python3 scripts/verification_form.py --filing "Form 10-K for fiscal 2025, accession 0001628280-26-008586"
    python3 scripts/verification_form.py --out "research/03 Drafts/Module 2 - Data verification.docx"
    python3 scripts/verification_form.py --workbook "workbook/QD [TICKER].xlsx"

It reads the one .xlsx in workbook/ (it never writes to it), lists every typed input in
the actual-year columns of the Detail Data tab, grouped by the sheet's own sections so
each statement in the filing is opened once, and writes a .docx using only the Python
standard library plus openpyxl (which the converter already needs).

Run it straight after scripts/fill_detail_data_from_edgar.py has written the actual columns.
It reads only typed cells, so it does not need the workbook saved in Excel first, as long as
the year headers Excel last saved agree with the date on FrontPage. Input rows the SEC data
could not fill are listed as blank, for the person to type from the filing.
"""
from __future__ import annotations

import datetime as dt
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is missing. Run: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "research" / "03 Drafts" / "Module 2 - Data verification.docx"

SECTIONS = [  # (pattern on the sheet's section header, heading on the form, where to look)
    (r"SUMMARY DATA", "Summary rows", "Income statement for revenue; the earnings per share note for diluted shares"),
    (r"INCOME STATEMENT", "Income statement", "Consolidated statements of income"),
    (r"BALANCE SHEET", "Balance sheet", "Consolidated balance sheets"),
    (r"CASH FLOW", "Cash flow statement", "Consolidated statements of cash flows"),
]
SKIP = r"VALUATION MULTIPLES|MARGINS"

try:  # the input rows of the annual layout, so a row the SEC data left blank still gets a line
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from fill_detail_data_from_edgar import ROWS as INPUT_ROWS
except Exception:  # pragma: no cover
    INPUT_ROWS = {}


def find_workbook() -> Path:
    files = [f for f in sorted((ROOT / "workbook").glob("*.xlsx")) if not f.name.startswith("~$")]
    if len(files) != 1:
        sys.exit(f"Expected exactly one .xlsx in workbook/, found {len(files)}.")
    return files[0]


def fmt(v) -> str:
    if v is None:
        return "(blank: type it from the filing)"
    return f"{v:,.1f}" if abs(v) >= 100 else f"{v:,.2f}"


# ---------- a minimal .docx writer (standard library only) ----------
def run(text: str, bold=False, size=20, color=None) -> str:
    props = "<w:rFonts w:ascii=\"Calibri\" w:hAnsi=\"Calibri\" w:cs=\"Calibri\"/>" + ("<w:b/>" if bold else "")
    if color:
        props += f'<w:color w:val="{color}"/>'
    props += f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
    return f'<w:r><w:rPr>{props}</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def para(text: str, bold=False, size=20, after=80, color=None) -> str:
    return f'<w:p><w:pPr><w:spacing w:before="0" w:after="{after}"/></w:pPr>{run(text, bold, size, color)}</w:p>'


def cell(text: str, width: int, bold=False, shade=None) -> str:
    tc = f'<w:tcW w:w="{width}" w:type="dxa"/>' + (f'<w:shd w:val="clear" w:color="auto" w:fill="{shade}"/>' if shade else "")
    return f'<w:tc><w:tcPr>{tc}</w:tcPr><w:p><w:pPr><w:spacing w:before="20" w:after="20"/></w:pPr>{run(text, bold, 18)}</w:p></w:tc>'


def table(header: list, rows: list, widths: list) -> str:
    b = '<w:top w:val="single" w:sz="4" w:space="0" w:color="808080"/><w:left w:val="single" w:sz="4" w:space="0" w:color="808080"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="808080"/><w:right w:val="single" w:sz="4" w:space="0" w:color="808080"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="808080"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    out = f'<w:tbl><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/><w:tblBorders>{b}</w:tblBorders><w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>'
    out += '<w:tr><w:trPr><w:tblHeader/></w:trPr>' + "".join(cell(h, w, True, "D9E2F3") for h, w in zip(header, widths)) + "</w:tr>"
    for r in rows:
        out += "<w:tr>" + "".join(cell(t, w) for t, w in zip(r, widths)) + "</w:tr>"
    return out + "</w:tbl>"


def write_docx(path: Path, body: str) -> None:
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    doc = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>{body}'
           '<w:sectPr><w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/><w:pgMar w:top="1080" w:right="1080" w:bottom="1080" w:left="1080" w:header="720" w:footer="720" w:gutter="0"/></w:sectPr>'
           "</w:body></w:document>")
    types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", types); z.writestr("_rels/.rels", rels); z.writestr("word/document.xml", doc)


def main() -> None:
    argv = sys.argv[1:]
    def opt(name, default=None):
        return argv[argv.index(name) + 1] if name in argv and argv.index(name) + 1 < len(argv) else default
    filing = opt("--filing", "")
    out = Path(opt("--out", str(DEFAULT_OUT)))
    if out.exists() and "--overwrite" not in argv:
        sys.exit(f"{out.name} already exists. It may hold checks someone typed. Rename it, or pass --overwrite to replace it.")
    path = Path(opt("--workbook")) if opt("--workbook") else find_workbook()
    if not path.exists():
        sys.exit(f"No such workbook: {path}")
    with zipfile.ZipFile(path) as z:
        unsaved = 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode("utf-8")
    vals = openpyxl.load_workbook(path, data_only=True)
    forms = openpyxl.load_workbook(path)
    if "Detail Data" not in vals.sheetnames:
        sys.exit("No 'Detail Data' tab in this workbook.")
    wv, wf = vals["Detail Data"], forms["Detail Data"]

    actual_cols = []
    for r in range(1, 6):
        found = [(c, str(wv.cell(r, c).value).strip()) for c in range(3, 20) if re.fullmatch(r"\d{4}A", str(wv.cell(r, c).value or "").strip())]
        if found:
            actual_cols = found; break
    if not actual_cols:
        sys.exit("Could not find actual-year headers (like 2025A) on Detail Data. Set the date on FrontPage, save in Excel, and run again.")
    if unsaved:
        # Typed cells are readable before Excel recalculates; the year headers are formula results
        # Excel saved earlier. Trust them only if they agree with the date now on FrontPage.
        d = forms["FrontPage"]["C4"].value if "FrontPage" in forms.sheetnames else None
        if isinstance(d, (int, float)):
            d = dt.datetime(1899, 12, 30) + dt.timedelta(days=d)
        want = [f"{d.year - 2}A", f"{d.year - 1}A"] if isinstance(d, (dt.date, dt.datetime)) else None
        if [lab for _, lab in actual_cols] != want:
            sys.exit(f"{path.name} has not been opened and saved in Excel since the last fill, and its year headers do not match the date on FrontPage. Open it in Excel, save, and run this again.")

    groups, current = [], None
    for r in range(1, wv.max_row + 1):
        label = wv.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        text = label.strip()
        if re.search(SKIP, text, re.I):
            current = None; continue
        hit = next((s for s in SECTIONS if re.search(s[0], text, re.I)), None)
        if hit:
            current = {"heading": hit[1], "where": hit[2], "rows": []}; groups.append(current); continue
        if current is None:
            continue
        typed = [(c, lab) for c, lab in actual_cols if not (isinstance(wf.cell(r, c).value, str) and str(wf.cell(r, c).value).startswith("="))]
        known = INPUT_ROWS.get(r)
        is_input = bool(known) and known[0] != "ebitda" and known[2] in text.lower()
        if not typed or not (is_input or any(isinstance(wv.cell(r, c).value, (int, float)) for c, _ in typed)):
            continue  # a formula row, or a row nobody types into
        derived = bool(known) and known[0] == "ebitda"
        for c, lab in typed:
            v = wv.cell(r, c).value
            if derived and not isinstance(v, (int, float)):
                continue
            name = text + (" (EBIT + D&A; not a line in the filing)" if derived else "")
            current["rows"].append([name, wf.cell(r, c).coordinate, lab, fmt(v if isinstance(v, (int, float)) else None), "", "", ""])
    groups = [g for g in groups if g["rows"]]
    total = sum(len(g["rows"]) for g in groups)
    if not total:
        sys.exit("No typed figures found in the actual-year columns. Fill the workbook first.")

    company = vals["FrontPage"]["A1"].value if "FrontPage" in vals.sheetnames else ""
    body = para("Module 2: Data verification", True, 32, 60, "1B2A4A")
    body += para(f"{company or '[Company]'}  ·  workbook {path.name}  ·  form made {dt.date.today().isoformat()}", False, 20, 60)
    body += para(f"Filing checked: {filing or '______________________________________________'}", False, 20, 60)
    body += para("Checked by: ____________________     Date: ______________", False, 20, 120)
    body += para("How to use this form. Every figure below was typed into the workbook's actual-year columns. Open the filing to the statement named above each table. "
                 "For each row, find the figure, type the page number, and type Y if it matches. If it does not match, type N and the filing's figure; if the workbook cell was blank, type the filing's figure. "
                 "The filing reports in thousands or millions; the workbook is in the units its row labels say. Save this file, then tell Claude you are done.", False, 18, 160)
    widths = [4300, 900, 900, 2000, 1300, 1300, 2980]
    header = ["Row on Detail Data", "Cell", "Year", "Value in workbook", "Page in filing", "Match (Y/N)", "Filing's figure, if different or blank"]
    for g in groups:
        body += para(f"{g['heading']}  ({len(g['rows'])} figures)", True, 24, 40, "1B2A4A")
        body += para(f"Where to look: {g['where']}", False, 18, 60)
        body += table(header, g["rows"], widths) + para("", after=160)
    body += para(f"Figures checked: ______ of {total}.     Figures corrected: ______.", True, 20, 60)
    write_docx(out, body)
    print(f"Wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}")
    print(f"  {total} figures in {len(groups)} blocks: " + ", ".join(f"{g['heading']} {len(g['rows'])}" for g in groups))
    print("  Open it in Word, fill in the page and Y or N for each row, save, and tell Claude you are done.")


if __name__ == "__main__":
    main()
